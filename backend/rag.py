"""Ingestion + retrieval on a local Chroma vector DB (persisted in data/chroma)."""
import csv
import re
from functools import lru_cache
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import EMBED_MODEL, RETRIEVAL_K, VECTOR_DIR

_splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=100)


@lru_cache(maxsize=1)
def get_store():
    from langchain_chroma import Chroma
    from langchain_huggingface import HuggingFaceEmbeddings
    return Chroma(
        collection_name="onboarding",
        embedding_function=HuggingFaceEmbeddings(model_name=EMBED_MODEL),
        persist_directory=str(VECTOR_DIR),
        collection_metadata={"hnsw:space": "cosine"},
    )


def _split_by_headings(text: str) -> list:
    sections, head, buf = [], "General", []
    for line in text.splitlines():
        if re.match(r"^#{1,6}\s+", line):
            if any(b.strip() for b in buf):
                sections.append((head, "\n".join(buf)))
            head, buf = line.lstrip("# ").strip(), []
        else:
            buf.append(line)
    if any(b.strip() for b in buf):
        sections.append((head, "\n".join(buf)))
    return sections


def read_sections(path: Path) -> list:
    """Return [(section_title, text)] for pdf / docx / csv / md / txt."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(f"Page {i + 1}", p.extract_text() or "") for i, p in enumerate(PdfReader(str(path)).pages)]
    if ext == ".docx":
        import docx
        lines = []
        for p in docx.Document(str(path)).paragraphs:
            if not p.text.strip():
                continue
            style = (p.style.name or "").lower()
            lines.append(("## " + p.text.strip()) if style.startswith("heading") or style == "title" else p.text.strip())
        return _split_by_headings("\n".join(lines))
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        return [("Training schedule", "\n".join("; ".join(f"{k}: {v}" for k, v in r.items()) for r in rows))]
    return _split_by_headings(path.read_text(encoding="utf-8", errors="ignore"))


def add_document(path: Path, filename: str, doc_type: str):
    """Chunk + embed. Returns (n_chunks, full_text_with_headings)."""
    sections = read_sections(path)
    docs = []
    for head, text in sections:
        for chunk in _splitter.split_text(text):
            docs.append(Document(page_content=f"[{head}] {chunk}",
                                 metadata={"source": filename, "section": head, "doc_type": doc_type}))
    if docs:
        get_store().add_documents(docs)
    return len(docs), "\n".join(f"## {h}\n{t}" for h, t in sections)


def remove_document(filename: str):
    try:
        get_store()._collection.delete(where={"source": filename})
    except Exception as exc:
        print(f"[rag] delete failed: {exc}")


def retrieve(query: str, k: int = RETRIEVAL_K) -> list:
    out = []
    for d, dist in get_store().similarity_search_with_score(query, k=k):
        out.append({
            "source": d.metadata.get("source", ""),
            "section": d.metadata.get("section", ""),
            "doc_type": d.metadata.get("doc_type", ""),
            "text": re.sub(r"^\[[^\]]*\]\s*", "", d.page_content),
            "score": round(max(0.0, 1 - float(dist)), 3),  # cosine distance -> similarity
        })
    return out
